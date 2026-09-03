import sqlite3
from datetime import datetime, timezone
from decimal import Decimal
from io import BytesIO

from openpyxl import load_workbook
from sqlalchemy import inspect, text

from ohmywealth.domain.models import Asset, AssetKind, MarketValue
from ohmywealth.infrastructure.database import create_database
from ohmywealth.infrastructure.database_catalog import (
    list_databases,
    remember_database,
    remembered_database,
    resolve_database,
)
from ohmywealth.infrastructure.demo import seed_demo_data
from ohmywealth.infrastructure.excel import (
    TABLE_ORDER,
    export_workbook,
    import_workbook,
)
from ohmywealth.infrastructure.repository import (
    SqliteMarketCache,
    SqliteWealthRepository,
)


def test_migration_and_market_cache_repository(tmp_path) -> None:
    database_path = tmp_path / "wealth.db"
    with sqlite3.connect(database_path) as connection:
        connection.executescript(
            """
            CREATE TABLE assets (
                id INTEGER PRIMARY KEY,
                symbol VARCHAR(64) NOT NULL UNIQUE,
                name VARCHAR(200) NOT NULL,
                kind VARCHAR(20) NOT NULL,
                currency VARCHAR(3) NOT NULL,
                exchange VARCHAR(80),
                archived BOOLEAN NOT NULL DEFAULT 0
            );
            CREATE TABLE schema_version (version INTEGER PRIMARY KEY);
            INSERT INTO schema_version (version) VALUES (1);
            """
        )
    engine = create_database(database_path)
    assert {
        "assets",
        "snapshots",
        "positions",
        "cached_prices",
        "cached_fx_rates",
        "settings",
    } <= set(inspect(engine).get_table_names())
    assert "nickname" in {
        column["name"] for column in inspect(engine).get_columns("assets")
    }
    with engine.connect() as connection:
        assert connection.scalar(text("select version from schema_version")) == 2

    cache = SqliteMarketCache(engine)
    quote = MarketValue(Decimal("9.99"), datetime.now(timezone.utc))
    cache.put_price("abc", quote)
    assert cache.get_price("ABC") == quote

    assert seed_demo_data(engine) is True
    assert seed_demo_data(engine) is False

    repository = SqliteWealthRepository(engine)
    asset = repository.add_asset(
        Asset(
            None,
            "ARCHIVE-TEST",
            "Archive test",
            AssetKind.CASH,
            "CNY",
            nickname="Rainy day",
        )
    )
    assert asset.id is not None
    assert asset.display_name == "Archive test (Rainy day)"
    repository.archive_asset(asset.id)
    assert all(item.id != asset.id for item in repository.list_assets())
    repository.unarchive_asset(asset.id)
    assert any(item.id == asset.id for item in repository.list_assets())

    workbook_bytes = export_workbook(engine, database_path)
    workbook = load_workbook(BytesIO(workbook_bytes))
    assert set(TABLE_ORDER) <= set(workbook.sheetnames)
    assets_sheet = workbook["assets"]
    assert assets_sheet["A1"].value == "id"
    assert assets_sheet.freeze_panes == "A2"
    assert len(assets_sheet.tables) == 1
    assert assets_sheet["A1"].fill.fgColor.rgb.endswith("087F70")
    workbook.close()

    imported_path = tmp_path / "imported.db"
    imported_engine = import_workbook(workbook_bytes, imported_path)
    imported_repository = SqliteWealthRepository(imported_engine)
    imported_asset = next(
        item
        for item in imported_repository.list_assets()
        if item.symbol == "ARCHIVE-TEST"
    )
    assert imported_asset.nickname == "Rainy day"
    assert len(imported_repository.list_snapshots()) == 6

    remember_database(imported_path)
    assert remembered_database(tmp_path) == imported_path
    assert imported_path in list_databases(imported_path)
    assert resolve_database(imported_path, imported_path.name) == imported_path
