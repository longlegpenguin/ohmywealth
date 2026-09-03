"""Excel backup export and validated import."""

from datetime import date, datetime, timezone
from decimal import Decimal, InvalidOperation
from io import BytesIO
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo
from sqlalchemy import select
from sqlalchemy.engine import Engine

from ohmywealth.domain.models import SUPPORTED_CURRENCIES, AssetKind
from ohmywealth.infrastructure.database import Base, create_database

TABLE_ORDER = (
    "assets",
    "snapshots",
    "positions",
    "cached_prices",
    "cached_fx_rates",
    "settings",
    "schema_version",
)
IMPORT_TABLES = TABLE_ORDER[:-1]
REQUIRED_SHEETS = ("assets", "snapshots", "positions")


def export_workbook(engine: Engine, database_path: Path) -> bytes:
    workbook = Workbook()
    overview = workbook.active
    overview.title = "README"
    overview.append(["Oh My Wealth database export"])
    overview.append(["Database", database_path.name])
    overview.append(
        ["Exported at", datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC")]
    )
    overview.append([])
    overview.append(
        [
            "Import notes",
            "Keep sheet and column names unchanged. IDs preserve table relationships.",
        ]
    )
    overview.column_dimensions["A"].width = 22
    overview.column_dimensions["B"].width = 72
    overview["A1"].font = Font(size=16, bold=True, color="FFFFFF")
    overview["A1"].fill = PatternFill("solid", fgColor="087F70")
    overview.merge_cells("A1:B1")
    overview.freeze_panes = "A2"
    overview.sheet_view.showGridLines = False

    with engine.connect() as connection:
        for table_name in TABLE_ORDER:
            table = Base.metadata.tables[table_name]
            primary_key = list(table.primary_key.columns)
            statement = select(table)
            if primary_key:
                statement = statement.order_by(*primary_key)
            rows = connection.execute(statement).mappings().all()
            sheet = workbook.create_sheet(table_name)
            headers = [column.name for column in table.columns]
            sheet.append(headers)
            for row in rows:
                values = [_excel_value(row[header]) for header in headers]
                sheet.append(values)
                for cell, value in zip(sheet[sheet.max_row], values, strict=True):
                    if isinstance(value, str) and value.startswith("="):
                        cell.data_type = "s"
            _format_sheet(sheet, table_name, len(rows), len(headers))

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def import_workbook(contents: bytes, target_path: Path) -> Engine:
    if target_path.exists():
        raise ValueError(f"Database already exists: {target_path.name}")
    try:
        workbook = load_workbook(BytesIO(contents), read_only=True, data_only=True)
    except Exception as error:
        raise ValueError(
            "The uploaded file is not a readable Excel workbook"
        ) from error

    try:
        missing = [name for name in REQUIRED_SHEETS if name not in workbook.sheetnames]
        if missing:
            raise ValueError(f"Workbook is missing sheets: {', '.join(missing)}")

        imported: dict[str, list[dict[str, Any]]] = {}
        for table_name in IMPORT_TABLES:
            if table_name not in workbook.sheetnames:
                imported[table_name] = []
                continue
            imported[table_name] = _read_sheet(workbook[table_name], table_name)
    finally:
        workbook.close()
    _validate_rows(imported)

    engine: Engine | None = None
    try:
        engine = create_database(target_path)
        with engine.begin() as connection:
            for table_name in IMPORT_TABLES:
                rows = imported[table_name]
                if rows:
                    connection.execute(Base.metadata.tables[table_name].insert(), rows)
        return engine
    except Exception as error:
        if engine is not None:
            engine.dispose()
        target_path.unlink(missing_ok=True)
        raise ValueError(f"Could not create database from workbook: {error}") from error


def _read_sheet(sheet: Any, table_name: str) -> list[dict[str, Any]]:
    rows = sheet.iter_rows(values_only=True)
    try:
        raw_headers = next(rows)
    except StopIteration as error:
        raise ValueError(f"Sheet {table_name} is empty") from error
    headers = [str(value).strip() if value is not None else "" for value in raw_headers]
    expected = [column.name for column in Base.metadata.tables[table_name].columns]
    if headers != expected:
        raise ValueError(f"Sheet {table_name} columns must be: {', '.join(expected)}")
    return [
        {
            header: _database_value(table_name, header, value)
            for header, value in zip(headers, values, strict=True)
        }
        for values in rows
        if any(value is not None for value in values)
    ]


def _validate_rows(tables: dict[str, list[dict[str, Any]]]) -> None:
    asset_ids = set()
    for asset in tables["assets"]:
        asset_ids.add(int(asset["id"]))
        if asset["kind"] not in {kind.value for kind in AssetKind}:
            raise ValueError(f"Unsupported asset kind: {asset['kind']}")
        if asset["currency"] not in SUPPORTED_CURRENCIES:
            raise ValueError(f"Unsupported currency: {asset['currency']}")
    snapshot_ids = {int(snapshot["id"]) for snapshot in tables["snapshots"]}
    for position in tables["positions"]:
        if int(position["asset_id"]) not in asset_ids:
            raise ValueError("A position refers to a missing asset")
        if int(position["snapshot_id"]) not in snapshot_ids:
            raise ValueError("A position refers to a missing snapshot")
        try:
            if Decimal(str(position["quantity"])) < 0:
                raise ValueError("Position quantities cannot be negative")
        except InvalidOperation as error:
            raise ValueError("A position contains an invalid quantity") from error


def _database_value(table_name: str, column: str, value: Any) -> Any:
    if value is None:
        return None
    if table_name == "snapshots" and column == "month":
        if isinstance(value, datetime):
            return value.date()
        if isinstance(value, date):
            return value
        return date.fromisoformat(str(value))
    if column == "observed_at":
        if isinstance(value, datetime):
            return value
        return datetime.fromisoformat(str(value))
    if table_name == "assets" and column == "archived":
        return bool(value)
    return value


def _excel_value(value: Any) -> Any:
    if isinstance(value, Decimal):
        return str(value)
    if isinstance(value, datetime) and value.tzinfo is not None:
        return value.replace(tzinfo=None)
    return value


def _format_sheet(sheet: Any, table_name: str, row_count: int, columns: int) -> None:
    header_fill = PatternFill("solid", fgColor="087F70")
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = Font(bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center")
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.sheet_view.showGridLines = False
    if row_count:
        table = Table(
            displayName=f"OhMyWealth{table_name.title().replace('_', '')}",
            ref=sheet.dimensions,
        )
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2",
            showRowStripes=True,
            showColumnStripes=False,
        )
        sheet.add_table(table)
    for index in range(1, columns + 1):
        values = [
            str(sheet.cell(row, index).value or "") for row in range(1, row_count + 2)
        ]
        sheet.column_dimensions[sheet.cell(1, index).column_letter].width = min(
            max(max(map(len, values), default=10) + 2, 12), 40
        )
